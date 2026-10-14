"""Import the supplied A/B v2 workbook without executing spreadsheet formulas."""
import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


def number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if np.isfinite(result) else None


def true(value):
    return value is True or value == 1 or str(value).strip().lower() == 'true'


def convert_workbook(source, output):
    source, output = Path(source), Path(output)
    formulas = load_workbook(source, data_only=False)
    cached = load_workbook(source, data_only=True)
    try:
        space = cached['Design_Space']
        params = {row[0].value: row[1].value for row in space.iter_rows() if row[0].value is not None}
        sea_selector, pcf_selector = params.get('SEA objective'), params.get('PCF objective')
        if sea_selector not in ('SEA_d', 'SEA_60') or pcf_selector not in ('PCF_preD', 'FirstPeak'):
            raise ValueError('Unsupported workbook objective selectors')
        sea_target = sea_selector + '_J_kg'
        pcf_target = 'PCF_preD_N' if pcf_selector == 'PCF_preD' else 'first_peak_force_N'
        constant_keys = ['L1_plus_L2_ref', 'N_cols', 'dx_ref', 'TargetWidth', 'WidthTolerance_pct', 'rho_AlSi10Mg']
        constants = {key: number(params.get(key)) for key in constant_keys}
        if any(v is None or v <= 0 for v in constants.values()):
            raise ValueError('Missing positive Design_Space geometry/QC constants')
        bounds = {}
        for row, name in [(32, 'r_L'), (33, 't'), (34, 'lambda_m')]:
            lo, hi = number(space.cell(row, 2).value), number(space.cell(row, 3).value)
            if lo is None or hi is None or lo >= hi:
                raise ValueError(f'Invalid workbook bounds: {name}')
            bounds[name] = (lo, hi)
        records, audit_rows, datasets = [], [], {}
        for sheet, label, id_column in [('Dataset_A_Controlled', 'A', 'Study_ID'), ('Dataset_B_ML', 'B', 'DOE_ID')]:
            ws, fs = cached[sheet], formulas[sheet]
            headers = {cell.value: cell.column for cell in fs[4] if cell.value is not None}
            required = {id_column, 'lambda_m', 'EA_60_J', 'SEA_60_J_g', 'EA_d_J', 'SEA_d_J_g',
                        'PCF_preD_N', 'FirstPeak_N', 'DensificationStrain', 'Mass_g', 'CoreVolume_mm3', 'JobStatus'}
            if label == 'B':
                required |= {'r_L', 't_mm', 'ActualWidth_mm', 'Max_ALLKE_ALLIE_pct', 'UseForML', 'GeometryValid'}
            if not required <= headers.keys():
                raise ValueError(f'{sheet} missing columns: {sorted(required - headers.keys())}')
            count = 0
            for row in range(5, ws.max_row + 1):
                def value(key):
                    return ws.cell(row, headers[key]).value if key in headers else None
                def formula(key):
                    return key in headers and fs.cell(row, headers[key]).data_type == 'f'
                def numeric(key):
                    return number(value(key))
                identity = value(id_column)
                if identity is None:
                    continue
                count += 1
                mass_g = numeric('Mass_g')
                if formula('Mass_g'):
                    volume = numeric('CoreVolume_mm3')
                    mass_g = None if volume is None else volume * constants['rho_AlSi10Mg']
                record = {'design_id': str(identity), 'dataset': label, 'workbook_sheet': sheet, 'workbook_row': row,
                          'data_origin': 'USER_WORKBOOK_FEM_RESULTS', 'lambda_m': numeric('lambda_m'),
                          'r_L': numeric('r_L') if label == 'B' else .5,
                          't': numeric('t_mm') if label == 'B' else numeric('t_final_CAD_mm'),
                          'mass_kg': None if mass_g is None else mass_g / 1000,
                          'PCF_preD_N': numeric('PCF_preD_N'), 'first_peak_force_N': numeric('FirstPeak_N'),
                          'epsilon_d': numeric('DensificationStrain'), 'JobStatus': value('JobStatus'),
                          'densification_status': 'not_available_from_summary'}
                for kind in ['60', 'd']:
                    energy = numeric(f'EA_{kind}_J')
                    sea = numeric(f'SEA_{kind}_J_g')
                    if formula(f'SEA_{kind}_J_g'):
                        sea = energy / mass_g if energy is not None and mass_g is not None and mass_g > 0 else None
                    record[f'EA_{kind}_J'] = energy
                    record[f'SEA_{kind}_J_kg'] = None if sea is None else sea * 1000
                record['SEA_Objective_J_kg'] = record[sea_target]
                record['PCF_Objective_N'] = record[pcf_target]
                reasons = []
                if label == 'A':
                    reasons.append('dataset_A_excluded_from_default_ML')
                else:
                    features_ok = all(record[k] is not None and bounds[k][0] <= record[k] <= bounds[k][1] for k in bounds)
                    if not features_ok:
                        reasons.append('invalid_or_missing_design_features')
                    geometry_ok = False
                    if features_ok:
                        total, cols = constants['L1_plus_L2_ref'], constants['N_cols']
                        shift = record['lambda_m'] * constants['dx_ref']
                        raw_l1 = record['r_L'] * total
                        pitch = total - 2 * shift
                        width = (cols - 1) * pitch + raw_l1
                        if pitch > 0 and width > 0:
                            scale = constants['TargetWidth'] / width
                            throat = raw_l1 * scale - 2 * shift * scale - 2 * record['t']
                            geometry_ok = record['lambda_m'] <= 0 or throat > 0
                    # Respect explicit manual rejection; recompute only this schema's known formula.
                    if not formula('GeometryValid'):
                        geometry_ok = geometry_ok and true(value('GeometryValid'))
                    if not geometry_ok:
                        reasons.append('nominal_geometry_invalid')
                    width = numeric('ActualWidth_mm')
                    width_error = None if width is None else 100 * (width - constants['TargetWidth']) / constants['TargetWidth']
                    record['width_error_pct_recomputed'] = width_error
                    if width_error is None or abs(width_error) > constants['WidthTolerance_pct']:
                        reasons.append('actual_width_missing_or_outside_tolerance')
                    kinetic = numeric('Max_ALLKE_ALLIE_pct')
                    if kinetic is None or not 0 <= kinetic <= 5:
                        reasons.append('kinetic_ratio_missing_or_outside_0_to_5_pct')
                    if value('JobStatus') != 'Valid':
                        reasons.append('JobStatus_not_Valid')
                    if not formula('UseForML') and not true(value('UseForML')):
                        reasons.append('UseForML_explicitly_disabled_or_missing')
                    if mass_g is None or mass_g <= 0:
                        reasons.append('positive_core_mass_missing')
                    if sea_selector == 'SEA_d' or pcf_selector == 'PCF_preD':
                        if record['epsilon_d'] is None or not 0 < record['epsilon_d'] < 1:
                            reasons.append('densification_strain_missing_or_invalid')
                    for target in (sea_target, pcf_target):
                        if record[target] is None or record[target] <= 0:
                            reasons.append(f'{target}_missing_or_nonpositive')
                record['qa_accepted'] = int(not reasons)
                record['UseForML_cached'] = value('UseForML')
                record['import_exclusion_reasons'] = ';'.join(reasons)
                records.append(record)
                audit_rows.append({'design_id': str(identity), 'sheet': sheet, 'excel_row': row,
                                   'eligible_for_selected_objectives': not reasons, 'exclusion_reasons': reasons})
            datasets[label] = count
        frame = pd.DataFrame(records)
        if frame[['dataset', 'design_id']].duplicated().any():
            raise ValueError('Duplicate design IDs within workbook dataset')
        config = {'seed': 42, 'test_size': .2, 'group_column': 'design_id',
                  'filters': {'dataset': 'B', 'qa_accepted': 1},
                  'features': [{'name': k, 'lower': lo, 'upper': hi, 'unit': 'mm' if k == 't' else 'dimensionless'} for k, (lo, hi) in bounds.items()],
                  'objectives': [{'name': sea_target, 'direction': 'max', 'unit': 'J/kg'},
                                 {'name': pcf_target, 'direction': 'min', 'unit': 'N'}],
                  'primary_model': 'xgboost', 'benchmark_models': ['xgboost', 'random_forest', 'svr'],
                  'cv_folds': 5, 'trees': 200, 'shap_samples': 100, 'population': 100, 'generations': 100}
        report = {'source_file': source.name, 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
                  'sheet_names': cached.sheetnames, 'dataset_rows': datasets,
                  'eligible_B_rows': int(frame.loc[frame.dataset == 'B', 'qa_accepted'].sum()),
                  'numeric_response_counts_B': {key: int(frame.loc[frame.dataset == 'B', key].notna().sum()) for key in
                      ['EA_60_J', 'SEA_60_J_kg', 'PCF_preD_N', 'epsilon_d', 'EA_d_J', 'SEA_d_J_kg']},
                  'objective_selectors': {'SEA': sea_selector, 'PCF': pcf_selector},
                  'units': {'source_SEA': 'J/g', 'output_SEA': 'J/kg', 'SEA_conversion_factor': 1000},
                  'qc_policy': 'Recompute known v2 nominal geometry, actual-width and KE<=5% predicates; require JobStatus=Valid, selected responses and positive core mass. Respect literal manual exclusions. This screens supplied summaries and does not establish physical FEM validity or raw-curve densification quality. ALLAE is retained in the source but has no threshold in workbook UseForML.',
                  'formula_policy': 'No formula evaluation. Recompute known mass/SEA/QC calculations from source fields; never use cached SEA/objective/UseForML formulas as FEM observations.',
                  'rows': audit_rows}
        output.mkdir(parents=True, exist_ok=False)
        frame.to_csv(output / 'fem_dataset.csv', index=False)
        for label in ('A', 'B'):
            # Preserve original columns/IDs instead of regenerating a different DOE.
            pd.read_excel(source, sheet_name='Dataset_A_Controlled' if label == 'A' else 'Dataset_B_ML', header=3).to_csv(
                output / f'original_dataset_{label}.csv', index=False)
        (output / 'workbook_config.json').write_text(json.dumps(config, indent=2))
        (output / 'workbook_audit.json').write_text(json.dumps(report, indent=2))
        return report
    finally:
        formulas.close()
        cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    result = convert_workbook(args.input, args.output)
    print(json.dumps({k: v for k, v in result.items() if k != 'rows'}, indent=2))
