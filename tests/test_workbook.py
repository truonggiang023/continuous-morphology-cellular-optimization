import json
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from cmco.workbook import convert_workbook
from cmco.pipeline import load_data


def make_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Design_Space'
    constants = {6: ('L1_plus_L2_ref', 24), 11: ('N_cols', 5), 14: ('rho_AlSi10Mg', .00267),
                 19: ('dx_ref', 4.029264821279436), 20: ('TargetWidth', 108),
                 29: ('WidthTolerance_pct', .5), 32: ('r_L', .4), 33: ('t_mm', .3),
                 34: ('lambda_m', -1), 38: ('SEA objective', 'SEA_d'), 39: ('PCF objective', 'PCF_preD')}
    for row, (key, value) in constants.items():
        ws.cell(row, 1, key); ws.cell(row, 2, value)
    for row, hi in [(32, .6), (33, .6), (34, 1)]:
        ws.cell(row, 3, hi)
    common = {'lambda_m': 0, 'EA_60_J': 5.34, 'SEA_60_J_g': '=AA5/Y5', 'EA_d_J': 2.67,
              'SEA_d_J_g': '=AF5/Y5', 'PCF_preD_N': 50, 'FirstPeak_N': 30,
              'DensificationStrain': .5, 'Mass_g': '=X5*Design_Space!B14',
              'CoreVolume_mm3': 1000, 'JobStatus': 'Valid'}
    for name, identity in [('Dataset_A_Controlled', 'Study_ID'), ('Dataset_B_ML', 'DOE_ID')]:
        sheet = wb.create_sheet(name)
        values = {identity: 'A_01' if identity == 'Study_ID' else 'B_001', **common}
        if identity == 'DOE_ID':
            values.update(r_L=.5, t_mm=.3, ActualWidth_mm=108, Max_ALLKE_ALLIE_pct=2,
                          GeometryValid='=TRUE', UseForML='=TRUE')
        for col, (key, value) in enumerate(values.items(), 1):
            sheet.cell(4, col, key); sheet.cell(5, col, value)
    wb.save(path)
    wb.close()


def test_workbook_formula_results_units_and_selectors(tmp_path):
    source = tmp_path / 'input.xlsx'
    make_workbook(source)
    report = convert_workbook(source, tmp_path / 'imported')
    assert report['eligible_B_rows'] == 1
    frame = pd.read_csv(tmp_path / 'imported' / 'fem_dataset.csv')
    b = frame[frame.dataset == 'B'].iloc[0]
    assert b.SEA_d_J_kg == pytest.approx(1000)
    assert b.SEA_60_J_kg == pytest.approx(2000)
    assert b.mass_kg == pytest.approx(.00267)
    assert b.densification_status == 'not_available_from_summary'
    config = json.loads((tmp_path / 'imported' / 'workbook_config.json').read_text())
    assert config['objectives'][0]['name'] == 'SEA_d_J_kg'
    assert config['objectives'][1]['name'] == 'PCF_preD_N'
    assert 'densification_status' not in config['filters']
    with pytest.raises(ValueError, match='at least 10'):
        load_data(tmp_path / 'imported' / 'fem_dataset.csv', config)


def test_manual_exclusion_and_missing_energy_not_filled(tmp_path):
    source = tmp_path / 'input.xlsx'
    make_workbook(source)
    wb = load_workbook(source)
    sheet = wb['Dataset_B_ML']
    headers = {cell.value: cell.column for cell in sheet[4]}
    sheet.cell(5, headers['UseForML'], 0)
    sheet.cell(5, headers['EA_d_J']).value = None
    wb.save(source); wb.close()
    report = convert_workbook(source, tmp_path / 'out')
    assert report['eligible_B_rows'] == 0
    row = report['rows'][1]
    assert 'SEA_d_J_kg_missing_or_nonpositive' in row['exclusion_reasons']
    assert 'UseForML_explicitly_disabled_or_missing' in row['exclusion_reasons']
    assert pd.isna(pd.read_csv(tmp_path / 'out' / 'fem_dataset.csv').SEA_d_J_kg.iloc[1])
